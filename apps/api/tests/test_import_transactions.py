"""Behavior tests for CSV/Excel transaction import."""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from tests.conftest import auth_header, make_token, seed_user
from tests.fakes import FakeClient

DEPT_ID = "dept-fin"
CAT_ID = "cat-sales"
CASH_ID = "cash-main"
PM_ID = "pm-transfer"
CSV_HEADER = ",".join([
    "transaction_date",
    "direction",
    "amount",
    "category_name",
    "department_code",
    "cash_account_name",
    "payment_method_name",
    "counterparty_name",
    "reference_no",
    "description",
])
CSV_REQUIRED_HEADER = ",".join(CSV_HEADER.split(",")[:7])


def _seed_refs(db: FakeClient) -> None:
    db.seed("departments", [{"id": DEPT_ID, "name": "Finance", "code": "FIN"}])
    db.seed(
        "cashflow_categories",
        [
            {"id": CAT_ID, "name": "Sales Income", "direction": "INFLOW"},
        ],
    )
    db.seed("cash_accounts", [{"id": CASH_ID, "name": "Main Bank"}])
    db.seed("payment_methods", [{"id": PM_ID, "name": "Bank Transfer"}])


def _csv(rows: list[str]) -> tuple[str, bytes, str]:
    content = "\n".join(rows).encode()
    return ("transactions.csv", content, "text/csv")


def test_finance_admin_imports_valid_csv_rows_as_draft(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "FINANCE_ADMIN")
    file = _csv(
        [
            CSV_HEADER,
            ",".join([
                "2026-07-01", "INFLOW", "1000000", "Sales Income", "FIN",
                "Main Bank", "Bank Transfer", "ACME", "REF-1", "Imported sale",
            ]),
        ]
    )

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": file},
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["total_rows"] == 1
    assert body["imported_count"] == 1
    assert body["failed_count"] == 0
    tx = fake_db.tables["cashflow_transactions"][0]
    assert tx["status"] == "DRAFT"
    assert tx["created_by"] == user_id
    assert tx["transaction_no"] == "INFLOW-202607-000001"
    assert any(
        log["action"] == "CREATE" for log in fake_db.tables["transaction_audit_logs"]
    )


def test_management_can_import_transactions(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "MANAGEMENT")
    file = _csv(
        [
            CSV_REQUIRED_HEADER,
            "2026-07-01,INFLOW,1000000,Sales Income,FIN,Main Bank,Bank Transfer",
        ]
    )

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": file},
    )

    assert resp.status_code == 200, resp.text
    assert resp.json()["imported_count"] == 1


def test_employee_cannot_import_transactions(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "EMPLOYEE", department_id=DEPT_ID)
    file = _csv(
        [
            CSV_REQUIRED_HEADER,
            "2026-07-01,INFLOW,1000000,Sales Income,FIN,Main Bank,Bank Transfer",
        ]
    )

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": file},
    )

    assert resp.status_code == 403


def test_import_partial_success_collects_row_errors(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "FINANCE_ADMIN")
    file = _csv(
        [
            CSV_HEADER,
            ",".join([
                "2026-07-01", "INFLOW", "1000000", "Sales Income", "FIN",
                "Main Bank", "Bank Transfer", "ACME", "REF-1", "Valid",
            ]),
            ",".join([
                "2026-07-02", "INFLOW", "500000", "Unknown", "FIN",
                "Main Bank", "Bank Transfer", "ACME", "REF-2", "Invalid",
            ]),
        ]
    )

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": file},
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["total_rows"] == 2
    assert body["imported_count"] == 1
    assert body["failed_count"] == 1
    assert body["errors"][0]["row_number"] == 3
    assert "category" in body["errors"][0]["message"].lower()
    assert len(fake_db.tables["cashflow_transactions"]) == 1


def test_import_rejects_more_than_500_rows(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "FINANCE_ADMIN")
    rows = [CSV_REQUIRED_HEADER] + [
        "2026-07-01,INFLOW,1000000,Sales Income,FIN,Main Bank,Bank Transfer"
        for _ in range(501)
    ]

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": _csv(rows)},
    )

    assert resp.status_code == 422


def test_import_rejects_ambiguous_cash_account_name(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    fake_db.seed("cash_accounts", [{"id": "cash-other", "name": "Main Bank"}])
    user_id = seed_user(fake_db, "FINANCE_ADMIN")
    file = _csv(
        [
            CSV_HEADER,
            ",".join([
                "2026-07-01", "INFLOW", "1000000", "Sales Income", "FIN",
                "Main Bank", "Bank Transfer", "ACME", "REF-1", "Ambiguous",
            ]),
        ]
    )

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": file},
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["imported_count"] == 0
    assert body["failed_count"] == 1
    assert "ambiguous" in body["errors"][0]["message"].lower()
    assert fake_db.tables.get("cashflow_transactions", []) == []


def test_import_rejects_oversized_file_before_parsing(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "FINANCE_ADMIN")
    content = b"x" * (10 * 1024 * 1024 + 1)

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": ("transactions.csv", content, "text/csv")},
    )

    assert resp.status_code == 422
    assert "10 MB" in resp.text


def test_import_rejects_invalid_csv_encoding(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "FINANCE_ADMIN")

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={"file": ("transactions.csv", b"\xff\xfe\xfa", "text/csv")},
    )

    assert resp.status_code == 422
    assert "Invalid CSV" in resp.text


def test_import_rejects_corrupt_xlsx(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "FINANCE_ADMIN")

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={
            "file": (
                "transactions.xlsx",
                b"not a spreadsheet",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert resp.status_code == 422
    assert "Invalid XLSX" in resp.text


def test_finance_admin_imports_xlsx_rows_as_draft(
    client: TestClient, fake_db: FakeClient
) -> None:
    _seed_refs(fake_db)
    user_id = seed_user(fake_db, "FINANCE_ADMIN")
    workbook = Workbook()
    ws = workbook.active
    ws.append(
        [
            "transaction_date",
            "direction",
            "amount",
            "category_name",
            "department_code",
            "cash_account_name",
            "payment_method_name",
            "counterparty_name",
            "reference_no",
            "description",
        ]
    )
    ws.append(
        [
            "2026-07-01",
            "INFLOW",
            1000000,
            "Sales Income",
            "FIN",
            "Main Bank",
            "Bank Transfer",
            "ACME",
            "REF-1",
            "Imported sale",
        ]
    )
    buf = BytesIO()
    workbook.save(buf)

    resp = client.post(
        "/api/import/transactions",
        headers=auth_header(make_token(user_id)),
        files={
            "file": (
                "transactions.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert resp.status_code == 200, resp.text
    assert resp.json()["imported_count"] == 1
    assert fake_db.tables["cashflow_transactions"][0]["status"] == "DRAFT"
